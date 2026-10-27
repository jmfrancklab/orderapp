"""Receipt exports preserve filtering, amounts, source pages and failure context."""
from io import BytesIO
from decimal import Decimal
import sqlite3

import pytest
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

import app as app_module
import receipt_reports as reports


@pytest.fixture(autouse=True)
def synthetic_ocr(monkeypatch):
    # Unit tests control recognized text; real rendering/OCR is tested separately.
    def recognize(page):
        if not hasattr(page, "cropbox"):
            return [{"text": page.extract_text(), "line": (0, 1, 10, 1)}]
        return reports.receipt_markup.text_spans(page)
    monkeypatch.setattr(reports.receipt_markup, "ocr_spans", recognize)


def source_pdf():
    buf = BytesIO()
    pdf = canvas.Canvas(buf, pagesize=(612, 792))
    pdf.drawString(40, 740, "Original supporting receipt")
    pdf.showPage()
    pdf.drawString(40, 740, "Supporting receipt second page")
    pdf.save()
    reader = PdfReader(BytesIO(buf.getvalue()))
    reader.pages[1].rotate(90)
    writer = PdfWriter()
    writer.append(reader)
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


@pytest.fixture
def export_client(tmp_path, monkeypatch):
    monkeypatch.setattr(app_module, "DB_PATH", str(tmp_path / "reports.db"))
    app_module.init_db()
    with sqlite3.connect(app_module.DB_PATH) as db:
        db.executemany("INSERT INTO projects (id, name) VALUES (?, ?)", [(1, "Alpha"), (2, "Beta")])
        db.executemany(
            "INSERT INTO invoices (id, nickname, created_by, created_at, receipt_url, invoice_url, reimbursement_status) VALUES (?, ?, ?, ?, ?, ?, ?)",
            [(11, "Mixed purchase", "buyer@lab.org", "2026-09-03", "https://files.example/receipt.pdf", "https://files.example/unused.pdf", "requires reimbursement"),
             (12, "Single purchase", "buyer@lab.org", "2026-09-02", "", "https://files.example/invoice.pdf", "reimbursed"),
             (13, "Private purchase", "private@lab.org", "2026-09-01", "", "", "requires reimbursement")])
        db.executemany(
            "INSERT INTO orders (description, user_email, invoice_id, project_id, cost, quantity, status) VALUES (?, ?, ?, ?, ?, ?, 'submitted')",
            [("Precision cable <shielded>", "buyer@lab.org", 11, 1, "12.35", 3),
             ("Adapter", "buyer@lab.org", 11, 2, "0.10", 3),
             ("Bench supply", "buyer@lab.org", 12, 1, "1000.10", 2)])
    downloads = []
    def download(url):
        downloads.append(url)
        return source_pdf()
    monkeypatch.setattr(reports, "download_document", download)
    with app_module.app.test_client() as client:
        with client.session_transaction() as session:
            session["email"] = "buyer@lab.org"
        yield client, downloads


def pdf_text(response):
    assert response.status_code == 200, response.get_data(as_text=True) if response.status_code != 200 else ""
    assert response.mimetype == "application/pdf"
    return "\n".join(page.extract_text() for page in PdfReader(BytesIO(response.data)).pages)


@pytest.mark.parametrize("grouping", ["receipt", "project"])
def test_export_amounts_sources_and_grouping(export_client, grouping):
    client, downloads = export_client
    response = client.get("/invoices/receipt-report", query_string={"group_by": grouping})
    text = pdf_text(response)
    assert "Receipt #11" in text and "Receipt #12" in text
    assert "Receipt #13" not in text
    assert "Precision cable <shielded>" in text
    for label in ("Quantity", "Unit cost", "Line total", "37.05", "0.30", "2,000.20"):
        assert label in text
    assert downloads == ["https://files.example/receipt.pdf", "https://files.example/invoice.pdf"]
    assert text.count("Original supporting receipt") == 3
    assert "Supporting receipt second page" in text
    assert "All purchases on this receipt are relevant to Alpha." in text
    assert "Only the following purchases on this receipt are relevant to Alpha." in text
    assert "this project" not in text
    assert "attachment" in response.headers["Content-Disposition"]


@pytest.mark.parametrize("grouping", ["receipt", "project"])
def test_nested_bookmarks_survive_supporting_document_merge(export_client, grouping):
    client, _ = export_client
    response = client.get("/invoices/receipt-report", query_string={"group_by": grouping})
    assert response.status_code == 200
    reader = PdfReader(BytesIO(response.data))
    outline = reader.outline
    assert reader.page_mode == "/UseOutlines"
    assert outline[0].title == "Report summary"
    if grouping == "project":
        assert outline[1].title == "Alpha"
        assert outline[2][0].title.startswith("Receipt #11 (Mixed purchase: $37.35)")
        assert outline[2][1][0].title.startswith("Supporting document")
        assert outline[3].title == "Beta"
        assert outline[4][0].title.startswith("Receipt #11")
    else:
        assert outline[1].title.startswith("Receipt #11")
        assert outline[2][0].title == "Alpha"
        assert outline[2][1][0].title.startswith("Supporting document")
        assert outline[2][2].title == "Beta"
        assert outline[2][3][0].title.startswith("Supporting document")

    def check_destinations(entries):
        for entry in entries:
            if isinstance(entry, list):
                check_destinations(entry)
            else:
                page_number = reader.get_destination_page_number(entry)
                assert page_number >= 0
                assert entry.title in reader.pages[page_number].extract_text()
    check_destinations(outline)


@pytest.mark.parametrize("selected,ids", [(["reimbursed"], [12]), (["requires reimbursement", "reimbursed"], [11, 12])])
def test_filter_matches_page_and_preserves_multiselect(export_client, selected, ids):
    client, _ = export_client
    text = pdf_text(client.get("/invoices/receipt-report", query_string={"filter_reimbursement_status": selected}))
    for number in (11, 12, 13):
        assert (f"Receipt #{number}" in text) == (number in ids)


def test_unexpected_generation_error_returns_log_reference(export_client, monkeypatch, caplog):
    client, _ = export_client
    def fail(*args):
        raise RuntimeError("PDF layout failed")
    monkeypatch.setattr(reports, "generate", fail)
    response = client.get("/invoices/receipt-report", query_string={"group_by": "project"})
    assert response.status_code == 500
    data = response.get_json()
    reference = data["message"].split("Error reference: ")[1].split(".")[0]
    assert reference in caplog.text
    assert "PDF layout failed" in caplog.text
    assert "grouping=project" in caplog.text
    assert "errors" not in data
    assert "PDF layout failed" not in data["message"]


def test_report_import_failure_returns_diagnostic_json(export_client, monkeypatch, caplog):
    import builtins
    client, _ = export_client
    original = builtins.__import__
    def importing(name, *args, **kwargs):
        if name == "receipt_reports":
            raise ImportError("Unavailable PDF library function")
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", importing)
    response = client.get("/invoices/receipt-report")
    assert response.status_code == 500
    assert "Error reference:" in response.get_json()["message"]
    assert "Unavailable PDF library function" in caplog.text


def test_missing_and_failed_files_identify_rows(export_client, monkeypatch):
    client, _ = export_client
    with sqlite3.connect(app_module.DB_PATH) as db:
        db.execute("UPDATE invoices SET receipt_url='', invoice_url='' WHERE id=11")
    monkeypatch.setattr(reports, "download_document", lambda url: b"not a document")
    response = client.get("/invoices/receipt-report")
    assert response.status_code == 400
    data = response.get_json()
    assert {item["invoice_id"] for item in data["errors"]} == {11, 12}
    assert "Receipt #11 (Mixed purchase)" in data["message"]
    assert "Filter to rows" in data["message"]
    assert "provide a receipt" in data["message"]
    first, second = data["errors"]
    assert (first["source_field"], first["source_url"], first["can_edit"]) == ("receipt_url", "", True)
    assert (second["source_field"], second["source_url"], second["can_edit"]) == ("invoice_url", "https://files.example/invoice.pdf", True)
    assert "receipt_url" not in second and "invoice_url" not in second


def test_bad_receipt_does_not_silently_fallback(export_client, monkeypatch):
    client, downloads = export_client
    monkeypatch.setattr(reports, "download_document", lambda url: (_ for _ in ()).throw(ValueError("Unreadable document.")))
    response = client.get("/invoices/receipt-report", query_string={"filter_reimbursement_status": "requires reimbursement"})
    assert response.status_code == 400
    assert len(response.get_json()["errors"]) == 1
    error = response.get_json()["errors"][0]
    assert error["source_field"] == "receipt_url"
    assert error["source_url"] == "https://files.example/receipt.pdf"
    assert not downloads


def test_fix_only_offending_link_then_retry(export_client, monkeypatch):
    client, _ = export_client
    def download(url):
        if url == "https://files.example/receipt.pdf":
            raise ValueError("Sharing permissions need fixing.")
        return source_pdf()
    monkeypatch.setattr(reports, "download_document", download)
    query = {"group_by": "project", "filter_reimbursement_status": "requires reimbursement"}
    response = client.get("/invoices/receipt-report", query_string=query)
    assert response.status_code == 400
    assert response.get_json()["errors"][0]["can_edit"]
    assert client.post("/api/invoices/11", json={"receipt_url": "https://files.example/shared.pdf"}).status_code == 200
    text = pdf_text(client.get("/invoices/receipt-report", query_string=query))
    assert "Receipt #11" in text and "Receipt #12" not in text
    with sqlite3.connect(app_module.DB_PATH) as db:
        assert db.execute("SELECT invoice_url, nickname FROM invoices WHERE id=11").fetchone() == ("https://files.example/unused.pdf", "Mixed purchase")


def test_excel_is_last_action_and_preserves_filters(export_client):
    from openpyxl import load_workbook
    import html
    import re
    client, _ = export_client
    query = [("filter_reimbursement_status", "reimbursed"), ("filter_reimbursement_status", "requires reimbursement")]
    page = client.get("/invoices", query_string=query).get_data(as_text=True)
    href = html.unescape(re.search(r'href="([^"]+)" class="mini export-btn"', page)[1])
    assert href.count("filter_reimbursement_status=") == 2
    assert page.index("Export to Excel") > page.index('data-receipt-report="project"')
    response = client.get(href)
    assert response.status_code == 200
    rows = list(load_workbook(BytesIO(response.data)).active.values)
    assert {row[0] for row in rows[1:]} == {11, 12}
    response = client.get("/export/invoices.xlsx?filter_reimbursement_status=reimbursed")
    rows = list(load_workbook(BytesIO(response.data)).active.values)
    assert [row[0] for row in rows[1:]] == [12]


def test_missing_cost_and_empty_results(export_client):
    client, _ = export_client
    with sqlite3.connect(app_module.DB_PATH) as db:
        db.execute("UPDATE orders SET cost='' WHERE invoice_id=11")
    response = client.get("/invoices/receipt-report")
    assert response.status_code == 400
    assert "Order #" in response.get_json()["message"]
    assert "invalid cost" in response.get_json()["message"]
    assert client.get("/invoices/receipt-report?filter_reimbursement_status=unknown").status_code == 400
    assert client.get("/invoices/receipt-report?group_by=invalid").status_code == 400


def test_shared_document_downloaded_once_and_unassigned(export_client):
    client, downloads = export_client
    with sqlite3.connect(app_module.DB_PATH) as db:
        db.execute("UPDATE invoices SET invoice_url='https://files.example/receipt.pdf' WHERE id=12")
        db.execute("UPDATE orders SET project_id=NULL WHERE invoice_id=12")
    assert "Unassigned project" in pdf_text(client.get("/invoices/receipt-report?group_by=project"))
    assert len(downloads) == 1


def test_images_and_placement_limits():
    for image_format in ("PNG", "JPEG"):
        buf = BytesIO()
        Image.new("RGB", (800, 1200), "white").save(buf, format=image_format)
        pages = reports.document_pages(buf.getvalue())
        placements = []
        flowable = reports.SupportingPage(pages[0], placements)
        assert flowable.width <= reports.CONTENT_WIDTH / 2
        assert flowable.height <= 600
    pages = reports.document_pages(source_pdf())
    assert float(pages[1].cropbox.width) == 792
    assert float(pages[1].cropbox.height) == 612


def test_export_requires_login_and_controls_disable_when_empty(export_client):
    client, _ = export_client
    page = client.get("/invoices?filter_reimbursement_status=unknown").get_data(as_text=True)
    assert 'data-receipt-report="receipt" disabled' in page
    assert 'data-receipt-report="project" disabled' in page
    with client.session_transaction() as session:
        session.clear()
    assert client.get("/invoices/receipt-report").status_code == 302


def test_long_tables_repeat_headers_and_split_long_descriptions():
    pages = reports.document_pages(source_pdf())
    items = [(Decimal("0.30"), f"Purchase {i}: " + "shielded cable " * 5, 3, Decimal("0.10"))
             for i in range(80)]
    items.append((Decimal("1.00"), "Long description " * 1200, 1, Decimal("1.00")))
    data = reports.build_pdf([{"id": 42, "groups": {(None, "Unassigned project"): items}, "pages": pages}], "receipt")
    reader = PdfReader(BytesIO(data))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert len(reader.pages) > 3
    assert text.count("Unit cost") > 3
    assert "25.00" in text
    assert "Original supporting receipt" in text


def test_public_download_rewrites_provider_and_rejects_private_redirects(monkeypatch):
    monkeypatch.setattr(reports.quotes, "classify_link", lambda url: "dropbox")
    monkeypatch.setattr(reports.socket, "getaddrinfo", lambda *args: [(2, 1, 6, "", ("8.8.8.8", 443))])
    requested = []
    class Response:
        is_redirect = True
        headers = {"Location": "http://127.0.0.1/private"}
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
    def get(self, url, **kwargs):
        requested.append(url)
        return Response()
    monkeypatch.setattr(reports.requests.Session, "get", get)
    def resolve(host, port):
        return [(2, 1, 6, "", ("127.0.0.1" if host == "127.0.0.1" else "8.8.8.8", port))]
    monkeypatch.setattr(reports.socket, "getaddrinfo", resolve)
    with pytest.raises(ValueError, match="public internet"):
        reports.download_document("https://www.dropbox.com/s/test/receipt.pdf?dl=0")
    assert requested == ["https://www.dropbox.com/s/test/receipt.pdf?dl=1"]


@pytest.mark.parametrize("url,query_key,query_value,cookie", [
    ("https://tenant-my.sharepoint.com/:b:/g/personal/buyer/example-token?e=share-key", "download", "1", "FedAuth"),
    ("https://www.dropbox.com/scl/fi/example-token/receipt.pdf?rlkey=share-key&dl=0", "dl", "1", "share_session"),
])
def test_share_download_preserves_cookies_and_tokens_through_redirects(monkeypatch, url, query_key, query_value, cookie):
    """Use real Requests cookie handling, with only its network transport replaced."""
    from email.message import Message
    from types import SimpleNamespace
    from urllib.parse import parse_qs, urlparse
    import requests

    monkeypatch.setattr(reports.socket, "getaddrinfo", lambda *args: [(2, 1, 6, "", ("8.8.8.8", 443))])
    pdf = source_pdf()
    seen = []

    def send(adapter, request, **kwargs):
        seen.append(request)
        response = requests.Response()
        response.request, response.url = request, request.url
        headers = Message()
        parsed = urlparse(request.url)
        if parsed.path != "/download/receipt.pdf":
            # Each download starts with a fresh anonymous session.
            assert "Cookie" not in request.headers
            query = parse_qs(parsed.query)
            assert query[query_key] == [query_value]
            assert query.get("rlkey", query.get("e")) == ["share-key"]
            response.status_code = 302
            response.headers["Location"] = "/download/receipt.pdf"
            headers["Set-Cookie"] = f"{cookie}=anonymous-share-access; Path=/; Secure; HttpOnly"
            response.headers["Set-Cookie"] = headers["Set-Cookie"]
            response.raw = BytesIO(b"")
        elif request.headers.get("Cookie") == f"{cookie}=anonymous-share-access":
            response.status_code = 200
            response.headers["Content-Type"] = "application/pdf"
            response.raw = BytesIO(pdf)
        else:
            # The previous implementation lost the cookie and reached this 403.
            response.status_code = 403
            response.raw = BytesIO(b"Access denied")
        response.raw._original_response = SimpleNamespace(msg=headers)
        return response

    monkeypatch.setattr(requests.adapters.HTTPAdapter, "send", send)
    for _ in range(2):
        downloaded = reports.download_document(url)
        assert downloaded == pdf
        assert len(reports.document_pages(downloaded)) == 2
    assert len(seen) == 4


@pytest.mark.parametrize('text,expected', [
    ('Order placed October 6, 2026\nReturn or replace items: Eligible through October 25, 2026', '2026-10-06'),
    ('Return or replace items: Eligible through October 25, 2026\nOrder SummaryOrder placed September 23, 2026', '2026-09-23'),
    ('Order Date\n9/11/26\nDelivers Sep 14', '2026-09-11'),
    ('Order Date: 0\n8/24/2026Bill To:', '2026-08-24'),
    ('Payment Date: August 28, 2026\nOrder Date: August 20, 2026', '2026-08-28'),
    ('Order Date Order Date 09/03/2026\nShipping Date: 09/07/2026', '2026-09-03'),
    ('Submitted 09/03/2026 10:27:25\n9/3/26, 11:27 AM DigiKey - Print', '2026-09-03'),
    ('INVOICE DATE\nSeptember 14, 2026 01:24 PM', '2026-09-14'),
    ('Placed September 14, 2026 $70.08\nEstimated Pick-up Date Today', '2026-09-14'),
    ('Date Fri 8/28/2026 3:33 PM', '2026-08-28'),
    ('Orde r Date: 14-Aug-2026\nShip Date: 15-Aug-2026', '2026-08-14'),
    ('Order Date: 2025-08-01', '2025-08-01'),
    ('Invoice Date: 09/03/2026\n123456789', '2026-09-03'),
    ('Document Date: 09/07/2026', '2026-09-07'),
    ('Docu ment Date: 09/07/2026', '2026-09-07'),
    ('Eligible through October 25, 2026', None),
    ('Order Date: 2/30/26', None),
    ('2026-08-01\n2026-08-02', '2026-08-01'),
    ('Order Date: 8/1/25\nOrder Date: 8/2/25', '2025-08-01'),
    ('Invoice DATE Order Status Comments 8/25/2026 Processing', '2026-08-25'),
    ('Date: Purchase details and reference number ABC123\n9/11/26', '2026-09-11'),
    ('Date\nOrder reference 123456789\n8/1/25', '2025-08-01'),
    ('Reference ABC 8/1/25 and 8/2/25', '2025-08-01'),
    ('An image-only receipt', None),
])
def test_transaction_date_uses_date_regex_then_numeric_fallback(text, expected):
    from types import SimpleNamespace
    result = reports.transaction_date([SimpleNamespace(extract_text=lambda: text)])
    assert (result.isoformat() if result else None) == expected


def test_date_uses_cached_ocr_for_scanned_receipts(monkeypatch):
    from types import SimpleNamespace
    calls = []
    def ocr(page):
        calls.append(page)
        return [{"text": "Payment Date: August 1, 2025", "line": (1, 2, 3, 2)}]
    monkeypatch.setattr(reports.receipt_markup, 'ocr_spans', ocr)
    page = SimpleNamespace(extract_text=lambda: '')
    geometry = [{}]
    assert reports.transaction_date([page], geometry).isoformat() == '2025-08-01'
    assert reports.transaction_date([page], geometry).isoformat() == '2025-08-01'
    assert len(calls) == 1


@pytest.mark.parametrize('grouping', ['receipt', 'project'])
def test_vendor_date_and_provider_are_in_both_report_modes(export_client, monkeypatch, grouping):
    client, _ = export_client
    with sqlite3.connect(app_module.DB_PATH) as db:
        db.execute("INSERT INTO vendors (id, name) VALUES (100, 'DigiKey')")
        db.execute('UPDATE orders SET vendor_id=100')
    def dated_pdf(url):
        buf = BytesIO()
        pdf = canvas.Canvas(buf)
        pdf.drawString(40, 740, 'Payment Date: August 1, 2025')
        pdf.save()
        return buf.getvalue()
    monkeypatch.setattr(reports, 'download_document', dated_pdf)
    response = client.get('/invoices/receipt-report', query_string={'group_by': grouping})
    text = pdf_text(response)
    # Invoice #11 spans two projects: its label must include both allocations.
    assert 'Receipt #11 (Mixed purchase: $37.35): DigiKey 8/1/25' in text
    assert 'Receipt #12 (Single purchase: $2,000.20): DigiKey 8/1/25' in text
    assert 'provided by: buyer@lab.org' in text
    labels = []
    for page in PdfReader(BytesIO(response.data)).pages:
        page.extract_text(visitor_text=lambda text, cm, tm, font, size: labels.append((text, size)))
    assert any('(Mixed purchase: $37.35)' in label and size == 9 for label, size in labels)
    assert 'Report summary' in text and 'OVERALL TOTAL:' in text
    # The repeated red tables and supporting pages must not double-count purchases.
    summary = PdfReader(BytesIO(response.data)).pages[0].extract_text()
    for amount in (('2,037.25', '0.30', '2,037.55') if grouping == 'project'
                   else ('37.35', '2,000.20', '2,037.55')):
        assert amount in summary
    if grouping == 'project':
        assert '#11 (Mixed purchase: $37.35)' in text
        assert '(Mixed purchase: $37.05)' not in text
        assert '(Mixed purchase: $0.30)' not in text
        assert text.count('OTHER PROJECTS TOTAL:') == 2
        assert 'PROJECT TOTAL:' in text
        assert 'Purchases for other projects' in text
    else:
        assert text.count('OTHER PROJECTS TOTAL:') == 2
        assert 'Purchases for other projects' in text
        assert 'RECEIPT TOTAL:' in text


@pytest.mark.parametrize('grouping', ['receipt', 'project'])
def test_other_project_names_and_totals_are_red_and_excluded(export_client, grouping):
    client, _ = export_client
    response = client.get('/invoices/receipt-report', query_string={'group_by': grouping})
    reader = PdfReader(BytesIO(response.data))
    seen_red = []
    for page in reader.pages:
        fill = (0, 0, 0)
        stack = []
        for operands, operator in page.get_contents().operations:
            if operator == b'q':
                stack.append(fill)
            elif operator == b'Q':
                fill = stack.pop() if stack else (0, 0, 0)
            elif operator == b'rg':
                fill = tuple(float(value) for value in operands)
            elif operator == b'Tj' and fill[0] > .6 and fill[1] == 0:
                seen_red.append(str(operands[0]))
    assert 'Beta' in seen_red and 'Alpha' in seen_red
    assert 'OTHER PROJECTS TOTAL:' in seen_red
    assert '0.30' in seen_red and '37.05' in seen_red
    # The single-project receipt is never listed as an excluded item.
    assert 'Bench supply' not in seen_red
    assert 'date not found' not in pdf_text(response)


def test_both_orders_use_identical_project_specific_support_and_cost_strikes(export_client, monkeypatch):
    client, _ = export_client
    def supporting_pdf(url):
        buf = BytesIO()
        pdf = canvas.Canvas(buf, pagesize=(612, 792))
        pdf.drawString(40, 700, 'Precision cable <shielded>')
        pdf.drawString(400, 700, '$37.05')
        pdf.drawString(40, 650, 'Adapter')
        pdf.drawString(400, 650, '$0.30')
        pdf.save()
        return buf.getvalue()
    monkeypatch.setattr(reports, 'download_document', supporting_pdf)
    original = reports.receipt_markup.strike_through
    calls = []
    def strike(pages, excluded, relevant, geometry, costs=None):
        lines, missing = original(pages, excluded, relevant, geometry, costs=costs)
        calls.append((excluded, relevant, costs, lines))
        assert not missing
        return lines, missing
    monkeypatch.setattr(reports.receipt_markup, 'strike_through', strike)
    rendered = []
    for grouping in ('project', 'receipt'):
        calls.clear()
        text = pdf_text(client.get('/invoices/receipt-report', query_string={'group_by': grouping}))
        assert text.count('Supporting document') == 3
        assert len(calls) == 2
        # Each allocation crosses out the other allocation's exact cost.
        for _, _, _, pages in calls:
            assert any(x0 >= 399 for page in pages for x0, _, _, _ in page)
        rendered.append(list(calls))
    assert rendered[0] == rendered[1]


@pytest.mark.parametrize('grouping', ['receipt', 'project'])
def test_unpositioned_amounts_download_report_instead_of_link_error(export_client, monkeypatch, grouping):
    client, _ = export_client
    def supporting_pdf(url):
        buf = BytesIO()
        pdf = canvas.Canvas(buf)
        pdf.drawString(40, 700, 'Precision cable $37.05')
        pdf.drawString(40, 650, 'Total $37.05')
        pdf.save()
        return buf.getvalue()
    monkeypatch.setattr(reports, 'download_document', supporting_pdf)
    original_spans = reports.receipt_markup.text_spans
    monkeypatch.setattr(reports.receipt_markup, 'text_spans', lambda page: original_spans(page)[:1])
    response = client.get('/invoices/receipt-report', query_string={'group_by': grouping})
    text = pdf_text(response)
    assert 'Receipt errors' in text
    assert 'no matching receipt item was found' in text
    assert 'Total $37.05' in text
    assert 'attachment' in response.headers['Content-Disposition']
